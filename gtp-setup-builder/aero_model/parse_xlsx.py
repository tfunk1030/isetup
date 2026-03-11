"""Parse individual aero map xlsx files into structured numpy arrays.

Two xlsx formats exist across the 5 GTP cars:

Labeled (BMW, Cadillac, Acura, Porsche):
  - Sheet "Raw Data" / "Data tables" / "AeroBalance"
  - Row 1: rear ride height values (columns B onward)
  - Col A: front ride height values (rows 2 onward)
  - Left block: DF balance (%)
  - Right block (after 3-4 empty cols): L/D ratio
  - Right block has its own front RH labels in a separate column

Unlabeled (Ferrari):
  - Sheet "AeroBalance"
  - No headers — row 1 is empty
  - Data starts at row 2, col 2
  - Same 51-row × 46-col grid, just without labels
  - Front RH assumed 25–75mm (1mm steps), rear RH assumed 5–50mm (1mm steps)
"""

from pathlib import Path

import numpy as np
import openpyxl


# Map from car directory name to the data sheet name
_SHEET_NAMES = {
    "BMWGTPAeroMap": "Raw Data",
    "CadallliacGTPAeromap": "Data tables",
    "Acuragtpaeromap": "AeroBalance",
    "Ferrarigtpaeromap": "AeroBalance",
    "Porschegtpaeromap": "AeroBalance",
}

# Map from directory name to canonical car name
CAR_NAMES = {
    "BMWGTPAeroMap": "bmw",
    "CadallliacGTPAeromap": "cadillac",
    "Acuragtpaeromap": "acura",
    "Ferrarigtpaeromap": "ferrari",
    "Porschegtpaeromap": "porsche",
}


def _find_data_sheet(wb: openpyxl.Workbook, car_dir: str) -> openpyxl.worksheet.worksheet.Worksheet:
    """Return the worksheet containing the raw aero data."""
    target = _SHEET_NAMES.get(car_dir)
    if target and target in wb.sheetnames:
        return wb[target]
    # Fallback: try common names
    for name in ("Raw Data", "Data tables", "Data Table", "AeroBalance"):
        if name in wb.sheetnames:
            sheet = wb[name]
            # Skip chart sheets
            if hasattr(sheet, "max_row"):
                return sheet
    raise ValueError(f"No data sheet found in workbook. Sheets: {wb.sheetnames}")


def _detect_format(ws) -> str:
    """Detect whether the sheet has labeled headers or is unlabeled (Ferrari)."""
    # Check if row 1 has numeric values in columns B+ (labeled format)
    for c in range(2, 10):
        v = ws.cell(row=1, column=c).value
        if v is not None and isinstance(v, (int, float)):
            return "labeled"
    return "unlabeled"


def _parse_labeled(ws) -> dict:
    """Parse a labeled-format sheet (BMW, Cadillac, Acura, Porsche)."""
    # Row 1, cols B onward: rear ride heights for balance block
    rear_rh_balance = []
    for c in range(2, ws.max_column + 1):
        v = ws.cell(row=1, column=c).value
        if v is None:
            break
        rear_rh_balance.append(float(v))

    n_rear = len(rear_rh_balance)

    # Col A, rows 2 onward: front ride heights
    front_rh = []
    for r in range(2, ws.max_row + 1):
        v = ws.cell(row=r, column=1).value
        if v is None:
            break
        front_rh.append(float(v))

    n_front = len(front_rh)

    # Balance data block: rows 2..(1+n_front), cols 2..(1+n_rear)
    balance = np.zeros((n_front, n_rear))
    for i in range(n_front):
        for j in range(n_rear):
            v = ws.cell(row=2 + i, column=2 + j).value
            balance[i, j] = float(v) if v is not None else np.nan

    # Find L/D block: scan rightward from end of balance block.
    # Must find a column where BOTH row 1 (header) and row 2 (data) have numeric values.
    # This skips the front-RH label column (which has data in rows but None in row 1).
    ld_start_col = None
    for c in range(2 + n_rear, ws.max_column + 1):
        v1 = ws.cell(row=1, column=c).value
        v2 = ws.cell(row=2, column=c).value
        if (v1 is not None and isinstance(v1, (int, float))
                and v2 is not None and isinstance(v2, (int, float))):
            ld_start_col = c
            break

    if ld_start_col is None:
        raise ValueError("Could not find L/D data block")

    # Read rear RH headers for L/D block (row 1, starting at ld_start_col)
    rear_rh_ld = []
    for c in range(ld_start_col, ws.max_column + 1):
        v = ws.cell(row=1, column=c).value
        if v is None:
            break
        rear_rh_ld.append(float(v))

    # L/D data
    ld = np.zeros((n_front, len(rear_rh_ld)))
    for i in range(n_front):
        for j in range(len(rear_rh_ld)):
            v = ws.cell(row=2 + i, column=ld_start_col + j).value
            ld[i, j] = float(v) if v is not None else np.nan

    return {
        "front_rh": np.array(front_rh),
        "rear_rh": np.array(rear_rh_balance),
        "balance": balance,
        "ld": ld,
    }


def _parse_unlabeled(ws) -> dict:
    """Parse an unlabeled-format sheet (Ferrari).

    No header row or label column. Data starts at row 2, col 2.
    Grid is 51 rows × 46 cols for each block.
    Assumed: front RH 25–75mm, rear RH 5–50mm.
    """
    # Count contiguous data columns in row 2 starting from col 2
    n_rear = 0
    for c in range(2, ws.max_column + 1):
        v = ws.cell(row=2, column=c).value
        if v is None:
            break
        n_rear += 1

    # Count contiguous data rows starting from row 2
    n_front = 0
    for r in range(2, ws.max_row + 1):
        v = ws.cell(row=r, column=2).value
        if v is None:
            break
        n_front += 1

    # Infer ride height axes (standard GTP grid)
    front_rh = np.arange(25, 25 + n_front, dtype=float)  # 25, 26, ..., 75
    rear_rh = np.arange(5, 5 + n_rear, dtype=float)  # 5, 6, ..., 50

    # Balance block
    balance = np.zeros((n_front, n_rear))
    for i in range(n_front):
        for j in range(n_rear):
            v = ws.cell(row=2 + i, column=2 + j).value
            balance[i, j] = float(v) if v is not None else np.nan

    # Find L/D block: scan rightward past the gap
    ld_start_col = None
    for c in range(2 + n_rear, ws.max_column + 1):
        v = ws.cell(row=2, column=c).value
        if v is not None and isinstance(v, (int, float)):
            ld_start_col = c
            break

    if ld_start_col is None:
        raise ValueError("Could not find L/D data block in unlabeled sheet")

    # Count L/D columns
    n_rear_ld = 0
    for c in range(ld_start_col, ws.max_column + 1):
        v = ws.cell(row=2, column=c).value
        if v is None:
            break
        n_rear_ld += 1

    ld = np.zeros((n_front, n_rear_ld))
    for i in range(n_front):
        for j in range(n_rear_ld):
            v = ws.cell(row=2 + i, column=ld_start_col + j).value
            ld[i, j] = float(v) if v is not None else np.nan

    return {
        "front_rh": front_rh,
        "rear_rh": rear_rh,
        "balance": balance,
        "ld": ld,
    }


def parse_aero_xlsx(filepath: str | Path) -> dict:
    """Parse a single aero map xlsx file.

    Returns dict with:
        car: str - canonical car name
        wing_angle: float - wing angle from filename
        front_rh: np.ndarray - front ride height axis (mm)
        rear_rh: np.ndarray - rear ride height axis (mm)
        balance: np.ndarray - DF balance (%) grid [front_rh × rear_rh]
        ld: np.ndarray - L/D ratio grid [front_rh × rear_rh]
    """
    filepath = Path(filepath)
    car_dir = filepath.parent.name
    car_name = CAR_NAMES.get(car_dir, car_dir)

    # Extract wing angle from filename
    wing_angle = _extract_wing_angle(filepath.name)

    wb = openpyxl.load_workbook(str(filepath), data_only=True, read_only=False)
    ws = _find_data_sheet(wb, car_dir)

    fmt = _detect_format(ws)
    if fmt == "labeled":
        data = _parse_labeled(ws)
    else:
        data = _parse_unlabeled(ws)

    wb.close()

    return {
        "car": car_name,
        "wing_angle": wing_angle,
        **data,
    }


def _extract_wing_angle(filename: str) -> float:
    """Extract wing angle from filename like 'Aero data 17 wing BMW LMDH.xlsx'."""
    # Pattern: number before "wing" in the filename
    parts = filename.lower().replace(".xlsx", "").split()
    for i, part in enumerate(parts):
        if part == "wing" and i > 0:
            try:
                return float(parts[i - 1])
            except ValueError:
                pass
    raise ValueError(f"Could not extract wing angle from filename: {filename}")
